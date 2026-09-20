#!/usr/bin/env python3
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def limpiar(texto):
    texto = re.sub(r"`([^`]*)`", r"\1", texto)
    texto = re.sub(r"\*\*([^*]*)\*\*", r"\1", texto)
    texto = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", texto)
    return texto.strip()


# Quita las barras de los extremos antes de partir: si no, salen dos celdas vacias por fila.
def celdas(linea):
    return [limpiar(c) for c in linea.strip().strip("|").split("|")]


def es_separador(linea):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", linea.strip()))


def tabla(doc, filas):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila[: len(filas[0])]):
            celda = t.cell(i, j)
            celda.text = valor
            for p in celda.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0
    doc.add_paragraph()


def convertir(origen, destino):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lineas = Path(origen).read_text(encoding="utf-8").splitlines()
    i = 0
    en_codigo = False
    codigo = []
    while i < len(lineas):
        linea = lineas[i]
        if linea.startswith("```"):
            if en_codigo:
                p = doc.add_paragraph("\n".join(codigo))
                p.paragraph_format.left_indent = Pt(18)
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)
                codigo = []
            en_codigo = not en_codigo
            i += 1
            continue
        if en_codigo:
            codigo.append(linea)
            i += 1
            continue
        if linea.startswith("|") and i + 1 < len(lineas) and es_separador(lineas[i + 1]):
            filas = [celdas(linea)]
            i += 2
            while i < len(lineas) and lineas[i].startswith("|"):
                filas.append(celdas(lineas[i]))
                i += 1
            tabla(doc, filas)
            continue
        if linea.startswith("#"):
            nivel = len(linea) - len(linea.lstrip("#"))
            doc.add_heading(limpiar(linea.lstrip("#")), min(nivel, 4))
        elif re.match(r"^\s*[-*] ", linea):
            doc.add_paragraph(limpiar(re.sub(r"^\s*[-*] ", "", linea)), style="List Bullet")
        elif re.match(r"^\s*\d+\. ", linea):
            doc.add_paragraph(limpiar(re.sub(r"^\s*\d+\. ", "", linea)), style="List Number")
        elif linea.strip() == "---":
            doc.add_paragraph()
        elif linea.strip():
            p = doc.add_paragraph(limpiar(linea))
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

    doc.save(destino)


if __name__ == "__main__":
    for ruta in sys.argv[1:]:
        salida = str(Path(ruta).with_suffix(".docx"))
        convertir(ruta, salida)
        print(salida)
